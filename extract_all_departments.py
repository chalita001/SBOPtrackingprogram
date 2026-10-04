import openpyxl
import json
import re

wb = openpyxl.load_workbook('SBOP.xlsx', data_only=True)

dept_mapping = {
    'Rev H(Molding MM)  ': {'code': 'MOLD', 'name': 'Molding / MM', 'sheet': 'Rev H(Molding MM)  '},
    'SBOP Rev H(Facility) ': {'code': 'FACILITY', 'name': 'Facility', 'sheet': 'SBOP Rev H(Facility) '},
    'SBOP Checklist Rev H (Assy) ': {'code': 'ASSY', 'name': 'Assembly', 'sheet': 'SBOP Checklist Rev H (Assy) '},
    'Rev H (WH)': {'code': 'WH', 'name': 'Warehouse', 'sheet': 'Rev H (WH)'},
    'SBOP  Rev H(Qc)': {'code': 'QC', 'name': 'Quality Control (QC)', 'sheet': 'SBOP  Rev H(Qc)'},
    'Rev H(Stamping) ': {'code': 'STAMPING', 'name': 'Stamping', 'sheet': 'Rev H(Stamping) '},
    'SBOP  Rev H(Tool)': {'code': 'TOOL', 'name': 'Tooling', 'sheet': 'SBOP  Rev H(Tool)'}
}

master_data = {}

for sheet_name, d_info in dept_mapping.items():
    if sheet_name not in wb.sheetnames:
        continue
    sheet = wb[sheet_name]
    
    current_layer = 'Layer 1'
    current_category = 'General'
    current_method = 'Observe'
    current_subcategory = ''
    
    items = []
    
    for r in range(1, sheet.max_row + 1):
        c1 = str(sheet.cell(r, 1).value or '').strip()
        c2 = str(sheet.cell(r, 2).value or '').strip()
        c3 = str(sheet.cell(r, 3).value or '').strip()
        c4 = str(sheet.cell(r, 4).value or '').strip()
        c5 = str(sheet.cell(r, 5).value or '').strip()
        
        row_str = f"{c1} {c2} {c3} {c4} {c5}"
        
        # Check Layer transitions
        if 'Layer 3' in row_str or 'เลเยอร์ 3' in row_str or 'Layer#3' in row_str:
            current_layer = 'Layer 3'
        elif 'Layer 2' in row_str or 'เลเยอร์ 2' in row_str or 'Layer#2' in row_str:
            current_layer = 'Layer 2'
        elif 'Layer 1' in row_str or 'เลเยอร์ 1' in row_str or 'Layer#1' in row_str:
            current_layer = 'Layer 1'
            
        # Check Category
        if c1 and ('I.' in c1 or 'II.' in c1 or 'III.' in c1 or 'IIII.' in c1 or 'IV.' in c1 or 'SBOP (Systems Audit)' in c1 or 'PPE' in c1 or 'ความเสี่ยง' in c1 or 'สภาพแวดล้อม' in c1):
            current_category = c1.replace('\n', ' ')
            
        # Check Method in C2
        if c2 and any(k in c2 for k in ['Check', 'สังเกต', 'Observe', 'ตรวจ']):
            current_method = c2.replace('\n', ' ')
            
        # Check Subcategory in C4 (e.g. A.ท่าทางการทำงาน, B.เครื่องจักร...)
        if c4 and re.match(r'^[A-Z]\.', c4):
            current_subcategory = c4.replace('\n', ' ')
            
        # Now detect questions in C3, C4, or C2
        # Layer 2 / Layer 3 often have questions in C3 or C4
        # Layer 1 usually has questions in C4
        candidate_text = ''
        item_order = None
        
        for val in [c3, c4, c2]:
            val_clean = val.strip()
            if not val_clean:
                continue
            # match numbered items like "1) ...", "1. ...", "1) "
            m = re.match(r'^(\d+)[\.\)]\s*(.*)', val_clean)
            if m:
                item_order = int(m.group(1))
                candidate_text = m.group(2).strip()
                if not candidate_text and c4 and val != c4:
                    candidate_text = c4.strip()
                break
                
        if candidate_text:
            items.append({
                'row_in_excel': r,
                'layer': current_layer,
                'category': current_category,
                'subcategory': current_subcategory,
                'method': current_method,
                'item_order': item_order,
                'question': candidate_text
            })
            
    master_data[d_info['code']] = {
        'code': d_info['code'],
        'name': d_info['name'],
        'sheet': sheet_name,
        'total_items': len(items),
        'items': items
    }

with open('master_sbop_questions.json', 'w', encoding='utf-8') as f:
    json.dump(master_data, f, ensure_ascii=False, indent=2)

print("Extracted all departments!")
for code, d in master_data.items():
    print(f"Dept {code}: {d['total_items']} items")
