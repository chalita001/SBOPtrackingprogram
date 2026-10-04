import openpyxl

wb = openpyxl.load_workbook('SBOP.xlsx', data_only=True)

with open('columns_analysis.txt', 'w', encoding='utf-8') as f:
    for name in wb.sheetnames:
        if name == 'History Record':
            continue
        sheet = wb[name]
        f.write(f"\n==================== SHEET: {name} (Rows: {sheet.max_row}, Cols: {sheet.max_column}) ====================\n")
        
        # Check rows 1 to 25
        for r in range(1, min(26, sheet.max_row + 1)):
            row_vals = []
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(r, c).value
                if val is not None:
                    row_vals.append(f"C{c}='{str(val).strip()}'")
            if row_vals:
                f.write(f"R{r:02d}: " + " | ".join(row_vals) + "\n")
                
        # Also check sample rows where results are filled (e.g., rows 12, 17, 24)
        f.write("--- Sample item rows ---\n")
        for r in range(12, min(35, sheet.max_row + 1)):
            cells = [f"C{c}='{sheet.cell(r, c).value}'" for c in range(1, sheet.max_column + 1) if sheet.cell(r, c).value is not None]
            if len(cells) > 1:
                f.write(f"R{r:02d}: " + " | ".join(cells) + "\n")

print("Analyzed columns successfully.")
