import openpyxl

wb = openpyxl.load_workbook('SBOP.xlsx', data_only=True)

with open('footers.txt', 'w', encoding='utf-8') as f:
    for name in wb.sheetnames:
        if name == 'History Record':
            continue
        sheet = wb[name]
        f.write(f"\n==================== FOOTER: {name} (Rows: {sheet.max_row}) ====================\n")
        start = max(1, sheet.max_row - 25)
        for r in range(start, sheet.max_row + 1):
            cells = [f"C{c}='{sheet.cell(r,c).value}'" for c in range(1, sheet.max_column + 1) if sheet.cell(r, c).value is not None]
            if cells:
                f.write(f"R{r:02d}: " + " | ".join(cells) + "\n")

print("Footers dumped successfully.")
